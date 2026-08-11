#!/usr/bin/env python3
"""Дописывает хвост рукописи «Сенсорные системы» в sensys|color_mllm.docx.

Переносит из STATYA_sensornye_sistemy.md незавершённые части (список литературы,
REFERENCES, подписи к рисункам) и чинит механические дефекты тела статьи.
Правки применяются кодом, а не руками (см. §6.5 плана). Запуск идемпотентен:
скрипт всегда стартует от sensys|color_mllm.BACKUP.docx.

    python scripts/finish_manuscript_docx.py
"""

import copy
import shutil
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "sensys|color_mllm.BACKUP.docx"
DST = ROOT / "sensys|color_mllm.docx"

applied, missed = [], []


# --------------------------------------------------------------------------- #
# Работа с run'ами: правка текста без потери форматирования
# --------------------------------------------------------------------------- #
def _replace_once(p, old, new, search_from):
    """Одна замена old->new начиная со смещения search_from.

    Совпадение ищется в склеенном тексте абзаца, поэтому корректно работает и
    когда подстрока разорвана между несколькими run'ами. Возвращает смещение,
    с которого продолжать поиск, или None, если совпадений больше нет. Поиск
    продолжается ПОСЛЕ вставленного текста, поэтому замены вида
    X -> X + дополнение не зацикливаются.
    """
    runs = p.runs
    full = "".join(r.text for r in runs)
    at = full.find(old, search_from)
    if at < 0:
        return None

    end = at + len(old)
    pos, done = 0, False
    for r in runs:
        r_start, r_end = pos, pos + len(r.text)
        pos = r_end
        if r_end <= at or r_start >= end:
            continue
        head = r.text[: max(0, at - r_start)]
        tail = r.text[max(0, end - r_start) :]
        r.text = head + ("" if done else new) + tail
        done = True
    return at + len(new)


def replace_in_paragraph(p, old, new, tag):
    """Заменяет ВСЕ вхождения old->new внутри абзаца."""
    n, offset = 0, 0
    while (offset := _replace_once(p, old, new, offset)) is not None:
        n += 1
    if n == 0:
        missed.append(f"{tag}: не найдено {old!r}")
        return False
    applied.append(tag if n == 1 else f"{tag} (×{n})")
    return True


def set_text(p, text, tag):
    """Полностью переписывает абзац, сохраняя форматирование первого run'а."""
    if not p.runs:
        missed.append(f"{tag}: пустой абзац")
        return False
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""
    applied.append(tag)
    return True


def clone_after(model_p, text, anchor_p):
    """Вставляет новый абзац после anchor_p, копируя формат model_p."""
    new_el = copy.deepcopy(model_p._p)
    anchor_p._p.addnext(new_el)
    new_p = docx.text.paragraph.Paragraph(new_el, model_p._parent)
    if new_p.runs:
        new_p.runs[0].text = text
        for r in new_p.runs[1:]:
            r.text = ""
    return new_p


def drop(p):
    p._p.getparent().remove(p._p)


# --------------------------------------------------------------------------- #
# Контент
# --------------------------------------------------------------------------- #

# --------------------------------------------------------------------------- #
# Раздел «Связанные работы» (перенос §1.1 из STATYA_sensornye_sistemy.md).
# Формулировки сверены с первоисточниками: см. комментарии.
# --------------------------------------------------------------------------- #
RELATED_WORK_HEADING = "Связанные работы"

RELATED_WORK = [
    # Abdou et al. 2021 (CoNLL, aclanthology.org/2021.conll-1.9): CIELAB, тёплые тона точнее.
    # Kawakita et al. 2024 (Sci Rep 14:15917): GWOT, 93 цвета, цвето-нормальные и цвето-атипичные
    #   наблюдатели, GPT-4 выравнивается заметно лучше GPT-3.5.
    # Marjieh et al. 2024 (Sci Rep 14:21445): шесть модальностей, цветовой круг; GPT-4, обученный
    #   совместно на зрении и языке, не даёт выигрыша именно в визуальной модальности.
    "Вопрос о том, воспроизводят ли языковые и мультимодальные модели перцептивную структуру цвета, "
    "изучается по нескольким направлениям. Пионерская работа (Abdou et al., 2021) показала, что уже "
    "в чисто текстовых языковых моделях эмбеддинги цветовых терминов структурно соответствуют "
    "перцептивно равномерному пространству CIELAB, причём тёплые тона выровнены точнее холодных. "
    "Последующие исследования усилили этот вывод на поведенческом уровне. Выравнивание без учителя "
    "методом Громова—Вассерштейна, применённое к оценкам сходства 93 цветов, выявило структурное "
    "соответствие матриц цветового сходства человека и больших языковых моделей, причём структура "
    "цвето-нормальных наблюдателей согласуется с GPT-4 заметно точнее, чем с GPT-3.5 "
    "(Kawakita et al., 2024). Парные оценки сходства, извлечённые из языковых моделей, воспроизводят "
    "психофизические структуры сразу в шести сенсорных модальностях, включая цветовой круг "
    "(Marjieh et al., 2024). Примечательно, что в последней работе совместное обучение GPT-4 на зрении "
    "и языке не давало выигрыша именно в визуальной модальности, — тогда как в настоящей работе "
    "представления визуального энкодера, напротив, точнее представлений языкового декодера на больших "
    "цветовых различиях.",

    # Ehab et al. 2026 (arXiv:2607.16540, 17.07.2026): >50 энкодеров, эллипсы Мак-Адама и пороги
    #   CIEDE2000, лучший mIoU < 0.25; self-supervised устойчиво лучше супервизорных, модели с
    #   языковым надзором — наиболее полярные.
    # Wickramanayaka, Oizumi 2025 (bioRxiv 10.64898/2025.12.10.693611): 16 моделей, те же 93 цвета,
    #   тонкое соответствие достигается ТОЛЬКО в парадигме CLIP.
    # Vision language models inherit human color perception (OpenReview 5rMHA5iSKV, 02.03.2026):
    #   Gemini 3 Flash и Qwen3-VL-8B-Instruct, >68 000 проб; послойное зондирование Qwen3-VL-8B —
    #   патч-эмбеддинги предпочитают sRGB (R²=0.97) метрике ΔE00 (R²=0.46).
    "Отдельная линия работ переходит от текста и поведенческих оценок к внутренним представлениям "
    "моделей зрения. Для более чем полусотни предобученных визуальных энкодеров показано, что "
    "расстояния между эмбеддингами, взятые без дополнительного обучения, лишь слабо согласуются "
    "с порогами цветоразличения человека: перекрытие с эллипсами Мак-Адама не превышает mIoU 0.25, "
    "причём энкодеры, обученные с самоконтролем (self-supervised), устойчиво превосходят обученные "
    "с учителем, а модели с языковым "
    "надзором дают наиболее полярные результаты (Ehab et al., 2026). Систематическое сравнение "
    "шестнадцати нейросетевых моделей с человеческими оценками сходства тех же 93 цветов приводит "
    "к согласующемуся выводу: тонкое структурное соответствие человеку достигается только в парадигме "
    "CLIP, тогда как модели, обученные с учителем или с самоконтролем, его не демонстрируют "
    "(Wickramanayaka, Oizumi, 2025). Послойный анализ мультимодальных моделей (Qwen3-VL-8B, "
    "Gemini 3 Flash) уточняет, где именно возникает перцептивная структура: патч-эмбеддинги входной "
    "проекции предпочитают линейную метрику sRGB (R² = 0.97) метрике ΔE₀₀ человека (R² = 0.46), тогда "
    "как поведение самой модели в задачах цветоразличения лучше всего объясняется именно ΔE₀₀ "
    "(Vision language models…, 2026).",

    "Эти результаты мотивируют избранный подход. Во-первых, перцептивную метрику следует извлекать "
    "не из входного пространства, а из глубоких представлений. Во-вторых, извлекать её нужно именно "
    "обучением: в «сыром» виде расстояния между эмбеддингами согласованы с психофизикой слабо, что "
    "подтверждается и на наших данных (см. Результаты).",

    "Со стороны колориметрии стандартом для оценки формул цветоразличения служит индекс STRESS "
    "(García et al., 2007), а современные равномерные пространства и формулы (семейство CAM16, "
    "CIEDE2000) калибруются и тестируются на объединённых наборах визуальных данных, в том числе "
    "COMBVD (Luo et al., 2023). На этом фоне вклад настоящей работы отличается от предшествующих "
    "в двух отношениях. Во-первых, вместо демонстрации корреляции или структурного соответствия "
    "представлений с эталонным пространством метрика обучается напрямую на первичных психофизических "
    "данных и оценивается той же процедурой, что и аналитические пространства, что переводит "
    "колориметрию из роли эталона в роль конкурента. Во-вторых, охватываются сразу два масштаба "
    "различий — большие (Манселл) и малые (COMBVD), — и показывается, что оба совмещаются одним "
    "линейным отображением.",
]

# Единый список источников. Все источники латиницей, поэтому по правилам журнала
# он приводится однократно под двумя заголовками. Порядок — алфавитный.
BIBLIOGRAPHY = [
    "Abdou M., Kulmizev A., Hershcovich D., Frank S., Pavlick E., Søgaard A. Can language models "
    "encode perceptual structure without grounding? A case study in color. Proceedings of the 25th "
    "Conference on Computational Natural Language Learning (CoNLL). 2021. P. 109–132. "
    "DOI: 10.18653/v1/2021.conll-1.9.",

    "Bai S., Chen K., Liu X., Wang J., Ge W., Song S., Dang K., Wang P., Wang S., Tang J., "
    "Zhong H., Zhu Y., Yang M., Li Z., Wan J., Wang P., Ding W., Fu Z., Xu Y., Ye J., Zhang X., "
    "Xie T., Cheng Z., Zhang H., Yang Z., Xu H., Lin J. Qwen2.5-VL technical report. arXiv preprint. "
    "2025. arXiv:2502.13923. DOI: 10.48550/arXiv.2502.13923.",

    "Ehab E., Hernández-Cámara P., Belal N., Malo J., Vazquez-Corral J., Gomez-Villa A. Do vision "
    "encoders exhibit human-like color thresholds? arXiv preprint. 2026. arXiv:2607.16540. "
    "DOI: 10.48550/arXiv.2607.16540.",

    "García P.A., Huertas R., Melgosa M., Cui G. Measurement of the relationship between perceived "
    "and computed color differences. J Opt Soc Am A. 2007. V. 24. № 7. P. 1823–1829. "
    "DOI: 10.1364/JOSAA.24.001823.",

    "Kawakita G., Zeleznikow-Johnston A., Tsuchiya N., Oizumi M. Gromov–Wasserstein unsupervised "
    "alignment reveals structural correspondences between the color similarity structures of humans "
    "and large language models. Scientific Reports. 2024. V. 14. № 1. Art. 15917. "
    "DOI: 10.1038/s41598-024-65604-1.",

    "Li C., Li Z., Wang Z., Xu Y., Luo M.R., Cui G., Melgosa M., Brill M.H., Pointer M. "
    "Comprehensive color solutions: CAM16, CAT16, and CAM16-UCS. Color Res Appl. 2017. V. 42. № 6. "
    "P. 703–718. DOI: 10.1002/col.22131.",

    "Liang Y., Li M., Fan C., Li Z., Nguyen D., Cobbina K., Bhardwaj S., Chen J., Liu F., Zhou T. "
    "ColorBench: Can VLMs see and understand the colorful world? A comprehensive benchmark for color "
    "perception, reasoning, and robustness. Advances in Neural Information Processing Systems. 2025. V. 38.",

    "Lobarev G., Kolomiets A., Karpenko S., Ershov E. Linear decoding of perceptual color spaces from "
    "multimodal large language models. London Imaging Meeting. 2026.",

    "Luo M.R., Cui G., Rigg B. The development of the CIE 2000 colour-difference formula: CIEDE2000. "
    "Color Res Appl. 2001. V. 26. № 5. P. 340–350. DOI: 10.1002/col.1049.",

    "Luo M.R., Xu Q., Pointer M., Melgosa M., Cui G., Li C., Xiao K., Huang M. A comprehensive test of "
    "colour-difference formulae and uniform colour spaces using available visual datasets. Color Res Appl. "
    "2023. V. 48. DOI: 10.1002/col.22844.",

    "Marjieh R., Sucholutsky I., van Rijn P., Jacoby N., Griffiths T.L. Large language models predict "
    "human sensory judgments across six modalities. Scientific Reports. 2024. V. 14. № 1. Art. 21445. "
    "DOI: 10.1038/s41598-024-72071-1.",

    "Robertson A.R. The CIE 1976 color-difference formulae. Color Research & Application. 1977. V. 2. "
    "№ 1. P. 7–11. DOI: 10.1002/j.1520-6378.1977.tb00104.x.",

    "Timofeev V., Usaev G., Seliugin M., Nikolaev D., Ershov E. No-reference error detection in color "
    "difference datasets: Application to Munsell data. IEEE Access. 2025. V. 13. P. 109322–109344. "
    "DOI: 10.1109/ACCESS.2025.3582053.",

    "Vision language models inherit human color perception. OpenReview preprint. 2026. "
    "URL: https://openreview.net/forum?id=5rMHA5iSKV (дата обращения: 10.08.2026).",

    "Wickramanayaka N.R., Oizumi M. Systematic comparison of color representations between humans and "
    "deep neural networks: towards predicting human color perception in a vast color space. bioRxiv. "
    "2025. DOI: 10.64898/2025.12.10.693611.",
]

# Перенумерация рисунков по порядку первого упоминания в тексте.
# Было: 1, 2, 3, 8, 9, 4, 5, 6, 7 — журнальное правило требует сквозного порядка.
# Ключ — прежний номер, значение — новый.
NEW_FROM_OLD = {1: 1, 2: 2, 3: 3, 8: 4, 9: 5, 4: 6, 5: 7, 6: 8, 7: 9}

# Соответствие «новый номер -> файл рисунка» (для вставки PDF в вёрстку):
#   1 fig0_graphical_abstract   4 fig8_gram_spectrum        7 fig5_map_spectrum
#   2 fig2_main_comparison      5 fig9_gram_spectrum_joint  8 fig6_testsize
#   3 fig3_per_group_munsell    6 fig4_manifold             9 fig7_error_heatmap

# Ссылки на рисунки в теле: (индекс абзаца, старая строка, новая строка).
# Замена двухфазная (через маркер \uE000 из Private Use Area), чтобы 8->4 не столкнулось с 4->6.
FIGURE_REFS = [
    (104, "рис. 8", "рис. \uE0004"),
    (105, "рис. 8в", "рис. \uE0004в"),
    (106, "рис. 8г", "рис. \uE0004г"),
    (106, "рис. 8а, б", "рис. \uE0004а, б"),
    (109, "Рис. 8", "Рис. \uE0004"),
    (113, "рис. 9в", "рис. \uE0005в"),
    (120, "рис. 9а, б", "рис. \uE0005а, б"),
    (123, "Рис. 9", "Рис. \uE0005"),
    (127, "рис. 4", "рис. \uE0006"),
    (129, "Рис. 4", "Рис. \uE0006"),
    (131, "рис. 5", "рис. \uE0007"),
    (134, "Рис. 5", "Рис. \uE0007"),
    (136, "рис. 6", "рис. \uE0008"),
    (138, "Рис. 6", "Рис. \uE0008"),
    (141, "рис. 7", "рис. \uE0009"),
    (145, "Рис. 7", "Рис. \uE0009"),
]

# Подписи к рисункам: ключ списка — ПРЕЖНИЙ номер (1..9), порядок и номера
# пересчитываются через NEW_FROM_OLD при записи.
CAPTIONS = [
    (
        "Рис. 1. Слева — цветовые центры Манселла в координатах CAM16-UCS (a′, b′, J′), рёбра соединяют "
        "соседей по единичному перцептивному шагу. Для выбранных пар цветовых центров показаны три отрезка "
        "в одном масштабе: единичный шаг по данным, предсказание CAM16-UCS и предсказание линейного "
        "отображения из представлений мультимодальной большой языковой модели Qwen2.5-VL-7B (оба с групповой "
        "калибровкой соответствующей оси). Справа — STRESS на надпороговом масштабе: линейное отображение "
        "представлений цвета МБЯМ превосходит всё семейство CAM16 и формулу CIEDE2000 — в 2.8 раза ближе "
        "к человеку, чем лучший вариант CAM16.",
        "Fig. 1. Left — Munsell color centers in CAM16-UCS coordinates (a′, b′, J′); edges connect neighbors "
        "separated by one unit perceptual step. For selected pairs of color centers, three segments are shown "
        "at the same scale: the unit step given by the data, the CAM16-UCS prediction, and the prediction of "
        "the linear map learned from the internal representations of the multimodal large language model "
        "Qwen2.5-VL-7B (both with group calibration of the corresponding axis). Right — STRESS on the "
        "suprathreshold scale: the linear map of MLLM color representations outperforms the entire CAM16 "
        "family and the CIEDE2000 formula, coming 2.8 times closer to the human data than the best CAM16 variant.",
    ),
    (
        "Рис. 2. Согласованность метрики с психофизикой на двух масштабах различий (столбцы — методы; "
        "штриховая линия — пол шума человека).",
        "Fig. 2. Agreement of the metric with human psychophysics on the two scales of color differences "
        "(bars — methods; dashed line — the human noise floor).",
    ),
    (
        "Рис. 3. Разложение по осям Манселла (тон/насыщенность/светлота): выигрыш отображения над CAM16 "
        "сосредоточен в тоне и насыщенности.",
        "Fig. 3. Decomposition along the Munsell axes (hue / chroma / lightness): the advantage of the learned "
        "map over CAM16 is concentrated in hue and chroma.",
    ),
    (
        "Рис. 4. PCA-спектр представлений: 95% дисперсии сосредоточено в 6 (визуальный энкодер) и 11 "
        "(языковой декодер) главных компонентах при номинальной размерности 3584.",
        "Fig. 4. PCA spectrum of the representations: 95% of the variance is contained in 6 (vision encoder) "
        "and 11 (language decoder) principal components at a nominal dimensionality of 3584.",
    ),
    (
        "Рис. 5. Структура обученного отображения A: сингулярный спектр, качество как функция ранга "
        "SVD-усечения (r ≈ 8 достигает уровня CAM16-LCD), значимость входных измерений и гистограмма весов.",
        "Fig. 5. Structure of the learned map A: singular spectrum, quality as a function of the SVD truncation "
        "rank (r ≈ 8 already matches CAM16-LCD), significance of the input dimensions, and the histogram of weights.",
    ),
    (
        "Рис. 6. Зависимость STRESS от размера тестовой выборки (Манселл и Leeds, m = 256, 3 сида): "
        "отображение превосходит CAM16-LCD даже при обучении на 10% цветов.",
        "Fig. 6. STRESS as a function of the test-set size (Munsell and Leeds, m = 256, 3 seeds): the map "
        "outperforms CAM16-LCD even when trained on 10% of the colors.",
    ),
    (
        "Рис. 7. Тепловая карта ошибки воспроизведения единичного шага Манселла по телу цветов "
        "(Value × Chroma — светлота × насыщенность; визуальный энкодер и языковой декодер). Мера — "
        "безразмерная величина |k_g·d̂ − 1| на тестовых парах, то есть по-парный вклад в STRESS "
        "(m = 256, 5-кратная кросс-валидация по цветам). В основном объёме цветового тела ошибка "
        "однородна и составляет в среднем 0.10 (визуальный энкодер) и 0.15 (языковой декодер) единицы "
        "шага; всплески до 0.41 и 0.64 соответственно сосредоточены у границы охвата sRGB при высокой "
        "насыщенности (C ≥ 18–20) и на светлотном пределе (V = 10).",
        "Fig. 7. Heat map of the error in reproducing the Munsell unit step over the color solid "
        "(Value × Chroma; vision encoder and language decoder). The measure is the dimensionless "
        "quantity |k_g·d̂ − 1| on the test pairs, i.e. the per-pair contribution to STRESS (m = 256, "
        "5-fold cross-validation over colors). Over the bulk of the color solid the error is uniform, "
        "averaging 0.10 (vision encoder) and 0.15 (language decoder) of a unit step; peaks of up to 0.41 "
        "and 0.64, respectively, are concentrated near the sRGB gamut boundary at high chroma "
        "(C ≥ 18–20) and at the lightness limit (V = 10).",
    ),
    (
        "Рис. 8. Спектр матрицы Грама G = AᵀA обученной метрики в зависимости от выходной размерности m "
        "(Манселл, group-k, 5-кратная CV по цветам): (а, б) собственные значения G при m = 32…3584 против "
        "шума инициализации — у языкового декодера над шумом ~5 собственных значений (<1% энергии метрики, "
        "превышение 1.5×), у визуального энкодера ~18 (≈5%, превышение 5.8×); (в) STRESS на тесте vs m — "
        "пол 0.152 ± 0.020 (LM, m ≈ 512–1024) против 0.096 ± 0.005 (VL, m ≈ 256–512); (г) функциональный "
        "ранг r_q (SVD-усечение с потерей ≤0.01 STRESS) — насыщение на ~170 (LM) против ~420 (VL) направлений "
        "при m = 3584. Языковой декодер хранит цвет в более сжатом представлении, чем визуальный энкодер.",
        "Fig. 8. Spectrum of the Gram matrix G = AᵀA of the learned metric as a function of the output "
        "dimensionality m (Munsell, group-k, 5-fold CV over colors): (a, b) eigenvalues of G for m = 32…3584 "
        "against the initialization-noise floor — the language decoder has ~5 eigenvalues above the floor "
        "(<1% of the metric energy, exceeding it by 1.5×), the vision encoder ~18 (≈5%, by 5.8×); (c) test "
        "STRESS vs m — a floor of 0.152 ± 0.020 (LM, m ≈ 512–1024) against 0.096 ± 0.005 (VL, m ≈ 256–512); "
        "(d) functional rank r_q (SVD truncation with a loss ≤0.01 STRESS) — saturation at ~170 (LM) against "
        "~420 (VL) directions at m = 3584. The language decoder stores color in a more compressed "
        "representation than the vision encoder.",
    ),
    (
        "Рис. 9. Совместное обучение на Манселле и Leeds (4 группы, групповой STRESS-лосс, разбиение по "
        "центрам в обоих наборах): (а, б) спектры матрицы Грама совместных отображений при разных m "
        "совпадают — быстрый спад до функционального ранга r_q ≈ 16 (LM) / 22 (VL), λ₃₂/λ₁ < 2%; "
        "(в) групповой STRESS выходит на плато ≈0.15 с m ≈ 32–64 без деградации до m = 3584 (пунктир — "
        "CAM16-LCD на тех же четырёх группах, 0.297); (г) r_q насыщается на ~20 у обоих уровней.",
        "Fig. 9. Joint training on Munsell and Leeds (4 groups, group STRESS loss, split by color centers in "
        "both datasets): (a, b) the Gram-matrix spectra of the joint maps coincide across m — a rapid decay "
        "to a functional rank r_q ≈ 16 (LM) / 22 (VL), λ₃₂/λ₁ < 2%; (c) the group STRESS reaches a plateau "
        "of ≈0.15 from m ≈ 32–64 with no degradation up to m = 3584 (dashed line — CAM16-LCD on the same four "
        "groups, 0.297); (d) r_q saturates at ~20 directions for both levels.",
    ),
]

# Механические правки тела: (индекс абзаца, что, на что, метка)
BODY_FIXES = [
    # — оборванные и битые ссылки —
    (50, "Qwen2.5-VL-7B [3].", "Qwen2.5-VL-7B (Bai et al., 2025).", "P50 висячая ссылка [3]"),
    # — источники, использованные без цитирования —
    (53, "COMBVD содержит психофизически", "База COMBVD (Luo et al., 2023) содержит психофизически",
     "P53 цитирование COMBVD"),
    (64, "STRESS (Standardized Residual Sum of Squares)",
     "STRESS (Standardized Residual Sum of Squares) (García et al., 2007)", "P64 цитирование STRESS"),
    # прямая ссылка на номер рисунка вперёд по тексту ломает сквозную нумерацию
    (107, ", ср. рис. 5)", ", см. ниже анализ структуры обученного отображения)",
     "P107 упреждающая ссылка на рисунок снята"),
    # — единообразие латинского «et al.» —
    (47, "et. al.", "et al.", "P47 et. al."),
    (48, "et. al.", "et al.", "P48 et. al."),
    # год издания формулы CIE 1976 — по списку литературы 1977
    (47, "(Robertson, 1976)", "(Robertson, 1977)", "P47 год Robertson (1976 -> 1977)"),
    (49, "(Lobarev, 2026)", "(Lobarev et al., 2026)", "P49 Lobarev et al."),
    (98, "(Lobarev, 2026)", "(Lobarev et al., 2026)", "P98 Lobarev et al."),
    (150, "(Lobarev, 2026)", "(Lobarev et al., 2026)", "P150 Lobarev et al."),
    # — перекрёстные ссылки на несуществующую нумерацию разделов —
    (53, "(см. раздел 5)", "(см. Обсуждение)", "P53 ссылка на раздел 5"),
    (114, "отмеченной в разделе 3.2", "отмеченной выше при групповой калибровке CAM16",
     "P114 ссылка на раздел 3.2"),
    (118, "(Значения раздела 3.3 —", "(Значения, приведённые выше при симметричном сравнении, —",
     "P118 ссылка на раздел 3.3"),
    (119, "выраженная в разделе 3.5 асимметрия", "отмеченная выше асимметрия", "P119 ссылка на раздел 3.5"),
    (131, "результат раздела 3.5 (отсутствие",
     "результат спектрального анализа матрицы Грама (отсутствие", "P131 ссылка на раздел 3.5"),
    (150, "метрики (раздел 3.5) отвечает", "метрики отвечает", "P150 ссылка на раздел 3.5"),
    (150, "цвета (раздел 3.9), но", "цвета, но", "P150 ссылка на раздел 3.9"),
    (151, "Совместная постановка (раздел 3.6) дополняет",
     "Совместное обучение на Манселле и Leeds дополняет", "P151 ссылка на раздел 3.6"),
    # — опечатки —
    (50, "над CAM16— в осях", "над CAM16 — в осях", "P50 пробел перед тире"),
    (82, "обучаемое оторбажение", "обучаемое отображение", "P82 опечатка «оторбажение»"),
    (112, "масштабом kg чем решается проблема разного мастшатба",
     "масштабом k_g, чем решается проблема разных масштабов", "P112 опечатка «мастшатба» + kg"),
    (124, "инвариантен к масштабу отображение", "инвариантен к масштабу отображения",
     "P124 согласование «отображение»"),
    (131, "отображения A (рис. 5). Её сингулярный", "отображения A (рис. 5). Его сингулярный",
     "P131 род «отображение»"),
    (100, "(языковой декодер.", "(языковой декодер).", "P100 незакрытая скобка"),
    (74, "евклидово пространство универсальное, для больших",
     "евклидово пространство — универсальное, для больших", "P74 пропущенное тире"),
    # — остатки markdown-разметки —
    (92, "По **светлоте**", "По светлоте", "P92 markdown **светлоте**"),
    (92, "в **тоне**", "в тоне", "P92 markdown **тоне**"),
    (92, "и **насыщенности**", "и насыщенности", "P92 markdown **насыщенности**"),
    # — единообразие подстрочных индексов в plain-text нотации —
    (55, "(обозначим eVL (vision layer embedding))", "(обозначим e_VL (vision layer embedding))", "P55 e_VL"),
    (55, "eLM (language model embedding)", "e_LM (language model embedding)", "P55 e_LM"),
    (59, "λ‖A‖²F", "λ‖A‖²_F", "P59 ‖A‖²_F"),
    (61, "где dij —", "где d_ij —", "P61 d_ij"),
    (106, "Функциональный ранг отображения rq (", "Функциональный ранг отображения r_q (", "P106 r_q (1)"),
    (106, "у визуального энкодера — rq ≈ 420", "у визуального энкодера — r_q ≈ 420", "P106 r_q (2)"),
    (109, "функциональный ранг rq (SVD", "функциональный ранг r_q (SVD", "P109 r_q"),
    (120, "функциональный ранг rq ≈ 16", "функциональный ранг r_q ≈ 16", "P120 r_q"),
    (123, "ранга rq ≈ 16", "ранга r_q ≈ 16", "P123 r_q"),
    (123, "(г) rq насыщается", "(г) r_q насыщается", "P123 r_q"),
    (124, "(10⁻⁵·‖A‖²F)", "(10⁻⁵·‖A‖²_F)", "P124 ‖A‖²_F"),
    (124, "размерности (mopt = 14–32)", "размерности (m_opt = 14–32)", "P124 m_opt"),
    (148, "(с весом KL)", "(с весом K_L)", "P148 K_L"),
]

# Абзацы, переписываемые целиком. Рис. 7 (после перенумерации — 9) строится по
# БЕЗРАЗМЕРНОЙ версии (единичные шаги Манселла), поэтому основной и контрольный
# абзацы меняются местами: метрическая карта — основная постановка, декодирование
# координат CAM16-LCD — контроль без рисунка. Числа: data/analysis/testsize_fig6_fig7.json
# (VL 0.103 / 0.408; LM 0.153 / 0.640) и data/analysis/fig7_error_dE.json (LM 0.546, VL 0.700).
PARA_REWRITES = [
    (141,
     "Пространственное распределение ошибок по осям Value×Chroma (светлота × насыщенность) "
     "оценивалось для самого метрического отображения (рис. 7). Мерой служит безразмерная ошибка "
     "воспроизведения единичного шага Манселла |k_g·d̂ − 1| — по-парный вклад в STRESS, вычисленный "
     "на тестовых парах при 5-кратной кросс-валидации по цветам. В основном объёме цветового тела "
     "ошибка мала и однородна (в среднем 0.10 единицы шага для визуального энкодера и 0.15 для "
     "языкового декодера), а локальные всплески до 0.41 и 0.64 соответственно сосредоточены у границы "
     "охвата sRGB при высокой насыщенности (C ≥ 18–20) и на светлотном пределе (V = 10) — в областях, "
     "слабо представленных и в данных Манселла, и в обучающем распределении модели. Картина "
     "воспроизводит наблюдение (Lobarev et al., 2026): внутренняя цветовая метрика МБЯМ наиболее "
     "надёжна в плотных областях распределения обучающих изображений.",
     "P141 основная постановка рис. 7 -> безразмерная (единичные шаги)"),

    (143,
     "Контрольная постановка — линейное декодирование координат CAM16-LCD из представлений "
     "(регрессия с регуляризацией Тихонова, та же кросс-валидация по цветам) — даёт зеркальную "
     "по слоям картину: "
     "средняя ошибка позиционирования составляет 0.55 ΔE для языкового декодера против 0.70 ΔE для "
     "визуального энкодера. Таким образом, на задаче «где находится цвет» точнее языковой декодер, "
     "а на метрической задаче «насколько два цвета различны» — визуальный энкодер. Противоречия здесь "
     "нет: задачи имеют разные эталоны и разную чувствительность. Декодирование координат — глобальная "
     "гладкая задача, почти насыщенная для обоих типов представлений (R² > 0.998, так что разница "
     "0.55 против 0.70 ΔE второго порядка), и оценивается она относительно конвенции CAM16-LCD. "
     "Воспроизведение единичных шагов — задача локальная: разность двух эмбеддингов удваивает "
     "по-точечный шум (у языкового декодера он выше из-за агрегации по длинной токенной "
     "последовательности с текстом), а качество определяется тонкой локальной геометрией "
     "представления, которая лучше сохранена в визуальном энкодере.",
     "P143 контрольная постановка (декодирование координат, без рисунка)"),
]


def main():
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    ps = d.paragraphs

    # 1. Абзацы, переписываемые целиком ------------------------------------
    for idx, text, tag in PARA_REWRITES:
        set_text(ps[idx], text, tag)

    # Подпись рис. 7 в теле — по безразмерной версии (совпадает с полосой подписей)
    set_text(ps[145], CAPTIONS[6][0], "P145 подпись рис. 7 -> безразмерная версия")

    # 2. Механические правки тела статьи -----------------------------------
    for idx, old, new, tag in BODY_FIXES:
        replace_in_paragraph(ps[idx], old, new, tag)

    # 3. Перенумерация рисунков по порядку первого упоминания ---------------
    for idx, old, new in FIGURE_REFS:
        replace_in_paragraph(ps[idx], old, new, f"P{idx} {old!r} -> {new.replace("\uE000", "")!r}")
    for p in ps:                                   # снятие временного маркера
        for r in p.runs:
            if "\uE000" in r.text:
                r.text = r.text.replace("\uE000", "")

    # 4. Раздел «Связанные работы» в конце Введения ------------------------
    head_model = docx.text.paragraph.Paragraph(copy.deepcopy(ps[54]._p), d)   # подзаголовок
    body_model = docx.text.paragraph.Paragraph(copy.deepcopy(ps[47]._p), d)   # абзац тела
    anchor = clone_after(head_model, RELATED_WORK_HEADING, ps[50])
    for para in RELATED_WORK:
        anchor = clone_after(body_model, para, anchor)
    applied.append(f"«Связанные работы»: заголовок + {len(RELATED_WORK)} абзаца после Введения")

    # 5. Абзац тела, ошибочно размеченный как Heading 1 --------------------
    ps[53].style = d.styles["Normal"]
    applied.append("P53 стиль Heading 1 -> Normal (абзац тела)")
    # Подзаголовок в стиле соседних (Normal + полужирный)
    ps[52].style = d.styles["Normal"]
    applied.append("P52 стиль Heading 1 -> Normal (как соседние подзаголовки)")
    # Единый стиль заголовков верхнего уровня
    for idx, name in [(44, "ВВЕДЕНИЕ"), (147, "ОБСУЖДЕНИЕ"), (153, "ЗАКЛЮЧЕНИЕ"),
                      (158, "СПИСОК ЛИТЕРАТУРЫ")]:
        ps[idx].style = d.styles["Heading 1"]
        applied.append(f"{name}: стиль -> Heading 1 (единообразие с ОПИСАНИЕ МЕТОДИКИ/РЕЗУЛЬТАТЫ)")

    # 6. Остаточные заглушки шаблона ---------------------------------------
    for idx, tag in [(156, "P156 заглушка «Текст раздела.»"), (142, "P142 случайный символ «ё»")]:
        drop(ps[idx])
        applied.append(tag)

    # 7. Список литературы + REFERENCES ------------------------------------
    # Все источники латиницей => по правилам журнала список даётся однократно
    # под двумя заголовками подряд.
    ref_model = copy.deepcopy(ps[165]._p)          # формат реальной библиозаписи
    refs_head_model = copy.deepcopy(ps[181]._p)    # заголовок REFERENCES

    for i in range(159, 198):                      # шаблон обоих списков
        drop(ps[i])

    anchor = ps[158]                               # «СПИСОК ЛИТЕРАТУРЫ»
    head_p = docx.text.paragraph.Paragraph(refs_head_model, anchor._parent)
    anchor._p.addnext(refs_head_model)
    head_p.style = d.styles["Heading 1"]
    anchor = head_p
    applied.append("REFERENCES: заголовок перенесён под СПИСОК ЛИТЕРАТУРЫ (единый список)")

    model_p = docx.text.paragraph.Paragraph(ref_model, d)
    for entry in BIBLIOGRAPHY:
        anchor = clone_after(model_p, entry, anchor)
    applied.append(f"Список литературы: {len(BIBLIOGRAPHY)} записей")

    # 8. Подписи к рисункам -------------------------------------------------
    ps = d.paragraphs
    cap_start = next(i for i, p in enumerate(ps) if p.text.strip() == "Подписи к рисункам")
    ru_model = copy.deepcopy(ps[cap_start + 2]._p)   # «Рис. 1. Подпись к первому рисунку.»
    blank_model = copy.deepcopy(ps[cap_start + 1]._p)

    for p in list(ps[cap_start + 1:]):
        drop(p)

    anchor = ps[cap_start]
    ru_m = docx.text.paragraph.Paragraph(ru_model, d)
    blank_m = docx.text.paragraph.Paragraph(blank_model, d)

    # порядок и номера — по карте перенумерации
    renumbered = []
    for old_no, (ru, en) in enumerate(CAPTIONS, start=1):
        new_no = NEW_FROM_OLD[old_no]
        ru = ru.replace(f"Рис. {old_no}.", f"Рис. {new_no}.", 1)
        en = en.replace(f"Fig. {old_no}.", f"Fig. {new_no}.", 1)
        renumbered.append((new_no, ru, en))
    renumbered.sort(key=lambda t: t[0])

    for _, ru, en in renumbered:
        anchor = clone_after(blank_m, "", anchor)
        anchor = clone_after(ru_m, ru, anchor)
        anchor = clone_after(ru_m, en, anchor)
    applied.append(
        f"Подписи к рисункам: {len(CAPTIONS)} × (рус. + англ.), нумерация по порядку упоминания")

    d.save(DST)

    print(f"Применено правок: {len(applied)}")
    for a in applied:
        print("  +", a)
    if missed:
        print(f"\nНЕ НАЙДЕНО: {len(missed)}")
        for m in missed:
            print("  !", m)
    return 1 if missed else 0


if __name__ == "__main__":
    sys.exit(main())
